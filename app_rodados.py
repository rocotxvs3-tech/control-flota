import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
import os

# ---------------------------------------------------------
# CONFIGURACIÓN DE PÁGINA
# ---------------------------------------------------------
st.set_page_config(page_title="Control de Rodados y Neumáticos", layout="wide", page_icon="🛞")

EXCEL_FILE = "CONTROL DE RODADOS XIMA SRL 2026.xlsx"

st.title("🛞 Módulo Independiente de Control de Rodados y Neumáticos")
st.markdown("Gestión de cambios, rotaciones e historial de cubiertas por unidad.")

# Verificar si el archivo Excel existe
if not os.path.exists(EXCEL_FILE):
    st.error(f"❌ No se encontró el archivo '{EXCEL_FILE}' en el directorio del proyecto. Por favor, asegúrate de subirlo.")
    st.stop()

# Cargar nombres de las hojas (unidades)
xls = pd.ExcelFile(EXCEL_FILE)
unidades_hojas = [h for h in xls.sheet_names if h != "Hoja1" and h != "Historial_Cambios"]

menu = st.sidebar.radio("Navegación:", [
    "📝 Registrar Cambio / Rotación de Neumático", 
    "📊 Ver Historial de Cambios",
    "🔍 Consultar Estado Actual de Unidad"
])

# ---------------------------------------------------------
# OPCIÓN 1: REGISTRAR CAMBIO O ROTACIÓN
# ---------------------------------------------------------
if menu == "📝 Registrar Cambio / Rotación de Neumático":
    st.subheader("🔄 Registrar Movimiento de Neumático")
    
    col1, col2 = st.columns(2)
    
    with col1:
        unidad_sel = st.selectbox("Seleccionar Unidad / Camión:", unidades_hojas, key="rod_unidad")
        posicion = st.selectbox("Posición del Neumático:", [
            "Eje Delantero Izquierdo",
            "Eje Delantero Derecho",
            "Eje Tracción Ext. Izquierdo",
            "Eje Tracción Int. Izquierdo",
            "Eje Tracción Int. Derecho",
            "Eje Tracción Ext. Derecho",
            "Eje Neumático / Balancín Izq.",
            "Eje Neumático / Balancín Der.",
            "Auxilio"
        ], key="rod_pos")
        
        km_actual = st.number_input("Kilometraje / Horas de la Unidad:", min_value=0, value=120000, step=500, key="rod_km")

    with col2:
        codigo_anterior = st.text_input("Código / Interno Neumático que se SACA:", placeholder="ej: INT 166", key="rod_ant")
        codigo_nuevo = st.text_input("Código / Interno Neumático que se COLOCA:", placeholder="ej: INT 180", key="rod_nue")
        marca_modelo = st.text_input("Marca y Modelo:", placeholder="ej: Fate / Pirelli / Bridgestone", key="rod_marca")
        
        motivo = st.selectbox("Motivo del Cambio:", [
            "Desgaste / Fin de Vida Útil",
            "Rotura / Pinchadura",
            "Recauche",
            "Rotación Programada entre Ejes",
            "Uso de Auxilio"
        ], key="rod_motivo")
        
        responsable = st.selectbox("Responsable / Mecánico:", [
            "Franco Luna",
            "Rocho",
            "Daniel",
            "Taller Externo"
        ], key="rod_resp")
        
        obs = st.text_area("Observaciones adicionales:", placeholder="Detalles de la maniobra...", key="rod_obs")

    st.divider()
    if st.button("💾 GUARDAR CAMBIO EN EL SISTEMA", type="primary", use_container_width=True):
        if not codigo_anterior or not codigo_nuevo:
            st.warning("⚠️ Debe especificar el código del neumático anterior y del nuevo.")
        else:
            try:
                wb = load_workbook(EXCEL_FILE)
                
                # 1. Asegurar que existe la hoja Historial_Cambios
                if "Historial_Cambios" not in wb.sheetnames:
                    ws_hist = wb.create_sheet("Historial_Cambios")
                    ws_hist.append([
                        "Fecha", "Unidad", "Posición", "Código Anterior", 
                        "Código Nuevo", "Marca / Modelo", "Kilometraje", 
                        "Motivo", "Responsable", "Observaciones"
                    ])
                else:
                    ws_hist = wb["Historial_Cambios"]
                
                # 2. Agregar registro al historial
                fecha_actual = datetime.now().strftime("%d/%m/%Y %H:%M")
                ws_hist.append([
                    fecha_actual,
                    unidad_sel,
                    posicion,
                    codigo_anterior,
                    codigo_nuevo,
                    marca_modelo,
                    km_actual,
                    motivo,
                    responsable,
                    obs
                ])
                
                wb.save(EXCEL_FILE)
                st.success(f"✅ ¡Cambio registrado correctamente para la unidad **{unidad_sel}** en la posición **{posicion}**!")
                st.info("💡 El historial se ha actualizado en la bitácora central.")

            except Exception as e:
                st.error(f"❌ Error al actualizar el archivo Excel: {e}")

# ---------------------------------------------------------
# OPCIÓN 2: VER HISTORIAL DE CAMBIOS
# ---------------------------------------------------------
elif menu == "📊 Ver Historial de Cambios":
    st.subheader("📜 Historial General de Movimientos de Neumáticos")
    
    if st.button("🔄 Actualizar Historial"):
        st.cache_data.clear()
        st.rerun()

    try:
        # Leer la hoja de historial directamente con pandas
        df_hist = pd.read_excel(EXCEL_FILE, sheet_name="Historial_Cambios")
        
        if df_hist.empty:
            st.info("Aún no hay registros en el historial de cambios.")
        else:
            # Filtro por unidad
            unidades_disponibles = ["Todas"] + list(df_hist["Unidad"].dropna().unique())
            filtro_hist_u = st.selectbox("Filtrar Historial por Unidad:", unidades_disponibles)
            
            df_mostrar = df_hist if filtro_hist_u == "Todas" else df_hist[df_hist["Unidad"] == filtro_hist_u]
            
            st.dataframe(df_mostrar.iloc[::-1], use_container_width=True)
            
            # Botón para descargar el historial actualizado
            with open(EXCEL_FILE, "rb") as file:
                st.download_button(
                    label="📥 Descargar Planilla Completa Actualizada (.xlsx)",
                    data=file,
                    file_name="CONTROL_DE_RODADOS_ACTUALIZADO.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

    except Exception:
        st.info("Aún no se ha generado la pestaña 'Historial_Cambios' en el archivo Excel. Realiza tu primer registro para crearla automáticamente.")

# ---------------------------------------------------------
# OPCIÓN 3: CONSULTAR ESTADO DE UNIDAD
# ---------------------------------------------------------
else:
    st.subheader("🔍 Consultar Estructura y Datos de Unidad")
    
    unidad_consulta = st.selectbox("Seleccione Unidad a Inspeccionar:", unidades_hojas, key="cons_unidad")
    
    try:
        df_unidad = pd.read_excel(EXCEL_FILE, sheet_name=unidad_consulta, header=None)
        st.markdown(f"### Vista de la Hoja: `{unidad_consulta}`")
        st.dataframe(df_unidad, use_container_width=True)
    except Exception as e:
        st.error(f"No se pudo cargar la hoja de la unidad: {e}")
